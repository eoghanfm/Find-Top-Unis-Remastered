import openpyxl
import json


def generate_excel():
    try:
        wb = openpyxl.load_workbook("uni thing.xlsx")
    except FileNotFoundError:
        wb = openpyxl.Workbook()


    sheet = wb.active

    #titles
    sheet["A1"] = "Rank"
    sheet["B1"] = "University"
    sheet["C1"] = "Course"
    sheet["D1"] = "Requirement"

    with open("universities.json", "r") as file:
        unis = json.load(file)

    uni_rank = 1
    for uni_name, courses in unis.items():

        ### FILTERS ###
        if uni_name == "Imperial College London" or uni_name == "University of Cambridge" or uni_name == "University of Oxford":
            uni_rank += 1
            continue

        for course in courses:
            row = sheet.max_row + 1
            #sheet[f"A{row}"] = row - 1  # Rank
            sheet[f"A{row}"] = uni_rank  # Rank

            sheet[f"B{row}"] = uni_name  # University
            sheet[f"C{row}"] = course["name"]  # Course
            sheet[f"D{row}"] = course["requirement"]  # Requirement

        uni_rank += 1

    #sheet.column_dimensions["A"].width = len(sheet["A"].value)
    #sheet.column_dimensions["B"].width = len(sheet["B"].value)
    #sheet.column_dimensions["C"].width = len(sheet["C"].value)
    #sheet.column_dimensions["D"].width = len(sheet["D"].value)


    sheet.column_dimensions["A"].width = max(len(str(cell.value or "")) for cell in sheet["A"]) * 0.9
    sheet.column_dimensions["B"].width = max(len(str(cell.value or "")) for cell in sheet["B"]) * 0.9
    sheet.column_dimensions["C"].width = max(len(str(cell.value or "")) for cell in sheet["C"]) * 0.9
    sheet.column_dimensions["D"].width = max(len(str(cell.value or "")) for cell in sheet["D"]) * 0.9

    # Add filters to headers
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"

    wb.save("uni thing.xlsx")

if __name__ == "__main__":
    generate_excel()